#!/usr/bin/env python3
"""
Población de cada municipio (DANE, proyecciones 2018-2042 del Censo 2018), resumida
para el tablero: total, cabecera y resto por año de elección, y estructura de
edades. Descarga los dos archivos oficiales del DANE:

  - PPED-AreaMun-2018-2042_VP.xlsx          población por área (4 MB)
  - PPED-AreaSexoEdadMun-2018-2042_VP.xlsx  población por área, sexo y edad simple (130 MB)

y escribe src/data/dane/municipios.json (el tablero no descarga nada en producción).
Uso:  pip install openpyxl && python3 scripts/snapshot-dane.py
"""
import json, os, re, sys, tempfile, urllib.request
import openpyxl

BASE = "https://www.dane.gov.co/files/censo2018/proyecciones-de-poblacion/Municipal/"
AREA = "PPED-AreaMun-2018-2042_VP.xlsx"
EDAD = "PPED-AreaSexoEdadMun-2018-2042_VP.xlsx"
# Años de elección con resultados en el tablero (2018 Congreso, 2022, 2023 territoriales, 2026).
ANIOS = [2018, 2022, 2023, 2026]
# Grupos de edad (desde, hasta): menores de 18, jóvenes, adultos y mayores.
GRUPOS = [(0, 17), (18, 29), (30, 44), (45, 59), (60, 200)]
OUT = os.path.join(os.path.dirname(__file__), "..", "src", "data", "dane", "municipios.json")


def bajar(nombre, tmp):
    ruta = os.path.join(tmp, nombre)
    if not os.path.exists(ruta):
        print("Descargando", nombre, "...", flush=True)
        req = urllib.request.Request(BASE + nombre, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=600) as r, open(ruta, "wb") as f:
            while chunk := r.read(1 << 20):
                f.write(chunk)
    return ruta


def hoja(ruta, nombre):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    return wb[nombre], wb["PPED"]


def main():
    tmp = os.environ.get("DANE_TMP") or tempfile.mkdtemp()
    municipios = {}

    # 1) Población por área.
    ws, meta = hoja(bajar(AREA, tmp), "PobMunicipalxÁrea")
    actualizado = next((str(r[0]) for r in meta.iter_rows(values_only=True) if r[0] and str(r[0]).startswith("Actualizado")), "")
    for dp, dpnom, mpio, nombre, anio, area, total in ws.iter_rows(min_row=9, max_col=7, values_only=True):
        if not mpio or anio not in ANIOS or area not in ("Total", "Cabecera Municipal", "Centros Poblados y Rural Disperso"):
            continue
        m = municipios.setdefault(str(mpio), {"nombre": nombre, "dep": str(dp), "departamento": dpnom, "pob": {}, "edad": {}})
        p = m["pob"].setdefault(str(anio), [0, 0, 0])  # total, cabecera, resto
        p[{"Total": 0, "Cabecera Municipal": 1, "Centros Poblados y Rural Disperso": 2}[area]] = int(total)
    print(len(municipios), "municipios con población")

    # 2) Estructura de edades (área "Total"), sumando hombres y mujeres por edad simple.
    ws, _ = hoja(bajar(EDAD, tmp), "PobMunicipalxÁreaSexoEdad")
    filas = ws.iter_rows(min_row=9, values_only=True)
    encabezado = next(filas)  # fila 9: nombres de columna ("Hombres 0 años" ...)
    columnas = []  # (índice, edad)
    for i, c in enumerate(encabezado):
        if isinstance(c, str):
            m = re.match(r"^(Hombres|Mujeres) (\d+)", c)
            if m:
                columnas.append((i, int(m.group(2))))
    print(len(columnas), "columnas de edad")
    for r in filas:
        mpio, anio, area = r[2], r[4], r[5]
        if not mpio or anio not in ANIOS or area != "Total" or str(mpio) not in municipios:
            continue
        grupos = [0] * len(GRUPOS)
        for i, edad in columnas:
            v = r[i]
            if v:
                for g, (d, h) in enumerate(GRUPOS):
                    if d <= edad <= h:
                        grupos[g] += int(v)
                        break
        municipios[str(mpio)]["edad"][str(anio)] = grupos

    salida = {
        "fuente": "DANE, proyecciones de población municipales 2018-2042 (base Censo 2018)",
        "actualizado": actualizado,
        "anios": ANIOS,
        "gruposEdad": ["0-17", "18-29", "30-44", "45-59", "60+"],
        "municipios": municipios,
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(salida, f, ensure_ascii=False, separators=(",", ":"))
    print("Escrito", os.path.normpath(OUT), f"({os.path.getsize(OUT) // 1024} KB)")


if __name__ == "__main__":
    sys.exit(main())
